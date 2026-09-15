import os,re
from pathlib import Path
from collections import Counter
import fitz
from docx import Document
from docx.enum.text import WD_BREAK, WD_LINE_SPACING
from docx.shared import Pt, Inches
FONT='Times New Roman'; SIZE=12
HEAD=re.compile(r'^(?:CHAPTER|BOOK|PART)\s+(?:[IVXLCDM]+|\d+|[A-Z]+)\b.*$',re.I); SPECIAL=re.compile(r'^(?:PROLOGUE|EPILOGUE|INTRODUCTION|PREFACE|FOREWORD|AFTERWORD)\s*$',re.I); DAY=re.compile(r'^(?:MONDAY|TUESDAY|WEDNESDAY|THURSDAY|FRIDAY|SATURDAY|SUNDAY)\.?\s*[-–—.:]*\s*$',re.I); PAGE=re.compile(r'^\s*(?:page\s+)?\d+\s*$',re.I)
def norm(s): return re.sub(r'[ \t]+',' ',s.replace('\u00a0',' ').replace('\u00ad','')).strip()
def heading(t): return bool(HEAD.match(t.strip()) or SPECIAL.match(t.strip()) or DAY.match(t.strip()))
def caps(t):
 t=t.strip(); a=[c for c in t if c.isalpha()]; return 3<=len(t)<=100 and bool(a) and all(c.isupper() for c in a)
def dialogue(t): return t.lstrip().startswith(('“','"','‘',"'")) or bool(re.match(r'^[—–-]\s+',t.lstrip()))
def repair(lines):
 out=[]; cur=''
 for raw in lines:
  line=raw.strip()
  if not line:
   if cur: out.append(cur.strip()); cur=''
   continue
  if PAGE.match(line): continue
  if not cur: cur=line; continue
  if cur.endswith('-') and re.search(r'[A-Za-z]-$',cur): cur=cur[:-1]+line
  else: cur+=' '+line
 if cur: out.append(cur.strip())
 return out
def extract(pdf):
 d=fitz.open(pdf); pages=[]
 for p in d:
  items=[]
  for b in p.get_text('blocks',sort=True):
   if len(b)<5: continue
   x0,y0,x1,y1,text=b[:5]; lines=text.replace('\r','').split('\n'); lines=[x.rstrip() for x in lines]
   while lines and not lines[0].strip(): lines.pop(0)
   while lines and not lines[-1].strip(): lines.pop()
   if lines: items.append({'x0':x0,'lines':lines,'raw':'\n'.join(lines)})
  pages.append(items)
 d.close(); return pages
def repeated(pages):
 c=Counter()
 for page in pages:
  for item in (page[:2]+page[-2:]):
   t=norm(item['raw'])
   if t and len(t)<=100:c[t]+=1
 return {t for t,n in c.items() if n>=max(3,int(len(pages)*.25))}
def paragraphs(pages):
 rep=repeated(pages); out=[]
 for pi,page in enumerate(pages,1):
  vals=[]
  for item in page:
   t=norm(item['raw'])
   if not t or t in rep or PAGE.match(t): continue
   for part in repair(item['lines']):
    if part: vals.append({'text':part,'x0':item['x0'],'page':pi,'heading':heading(part) or caps(part),'dialogue':dialogue(part)})
  for i,v in enumerate(vals): v['page_end']=i==len(vals)-1
  out.extend(vals)
 return out
def write_docx(items,path,preserve=True,smart=True):
 d=Document(); s=d.sections[0]
 for a in ('top_margin','bottom_margin','left_margin','right_margin'): setattr(s,a,Inches(1))
 n=d.styles['Normal']; n.font.name=FONT; n.font.size=Pt(SIZE); n.paragraph_format.line_spacing_rule=WD_LINE_SPACING.SINGLE; n.paragraph_format.space_before=Pt(0); n.paragraph_format.space_after=Pt(0)
 for i,x in enumerate(items):
  p=d.add_paragraph(); p.style=n; p.paragraph_format.line_spacing=1; p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0); r=p.add_run(x['text']); r.font.name=FONT; r.font.size=Pt(SIZE)
  if x['heading']: p.alignment=1; r.bold=True
  elif x['dialogue']: p.paragraph_format.left_indent=Inches(.25)
  elif x['x0']>45:
   ind=min(max((x['x0']-45)/72,0),1)
   if ind>=.08:p.paragraph_format.first_line_indent=Inches(ind)
  if preserve and x['page_end'] and i<len(items)-1 and items[i+1]['page']!=x['page']:
   y=items[i+1]; ok=True
   if smart: ok=y['heading'] or y['dialogue'] or bool(y['text'][:1].isupper())
   if ok:r.add_break(WD_BREAK.PAGE)
 d.save(path)
def convert_pdf_to_docx(pdf,workdir,preserve_page_breaks=True,smart_page_breaks=True):
 items=paragraphs(extract(pdf))
 if not items: raise ValueError('No selectable text was found. This may be a scanned/image-only PDF.')
 stem=Path(pdf).stem; docx=os.path.join(workdir,stem+'.docx')
 write_docx(items,docx,preserve_page_breaks,smart_page_breaks)
 clean=os.path.join(workdir,stem+'_clean.txt'); review=os.path.join(workdir,stem+'_review.txt')
 with open(clean,'w',encoding='utf8') as f:f.write('\n\n'.join(x['text'] for x in items)+'\n')
 with open(review,'w',encoding='utf8') as f:
  f.write('PDF → Word REVIEW REPORT\n========================\n\nThis report flags items for human review; it does not silently rewrite the manuscript.\n\n')
  found=[]
  for x in items:
   for w in re.findall(r'\b[A-Za-z][A-Za-z\'’-]{2,}\b',x['text']):
    if len(w)>28: found.append((w,x['page'],x['text']))
  f.write('Unusually long words:\n\n' if found else 'No unusually long words were detected.\n')
  seen=set()
  for w,p,t in found:
   if (w.lower(),p) not in seen: seen.add((w.lower(),p)); f.write(f'- Page {p}: {w}\n  {t}\n')
 return {'docx_path':docx,'docx_name':stem+'.docx'}
